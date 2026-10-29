# -*- coding: utf-8 -*-
"""生成给用户填写的「HS 对照表」。
数据源：master_seed/bom_full.json（飞书 BOM 清洗结果）
输出：HS对照表_待填.xlsx
列：ERP商品编码 | 品名(中) | HS编码(请填) | 型号(请核对) | 品牌 | 单位 | 单价USD | 单件毛重KG | 单件净重KG
"""
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
BOM = os.path.join(BASE, 'master_seed', 'bom_full.json')
OUT = os.path.join(BASE, 'HS对照表_待填.xlsx')

# 已确认为真的 HS（其余留空由用户填）
KNOWN_HS = {'TJ00033': '9031499000'}

HEADERS = ['ERP商品编码', '品名(中)', 'HS编码(请填)', '型号(请核对)', '品牌',
           '单位', '单价USD', '单件毛重KG', '单件净重KG']


def guess_unit(name, model, cat):
    s = (name or '') + (model or '') + (cat or '')
    if any(k in s for k in ('主机', '模组', '扫描仪')):
        return '套'
    return '个'


def main():
    recs = json.load(open(BOM, encoding='utf-8'))
    rows = []
    for r in recs:
        if not r.get('salable'):
            continue
        erp = r.get('erp', '')
        if not erp:
            continue
        rows.append([
            erp,
            r.get('name', ''),
            KNOWN_HS.get(erp, ''),
            r.get('model', ''),
            r.get('brand', ''),
            guess_unit(r.get('name'), r.get('model'), r.get('cat')),
            '', '', '',
        ])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'HS对照表'
    ws.append(HEADERS)

    hfill = PatternFill('solid', fgColor='2F5597')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF', size=11)
        c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row in rows:
        ws.append(row)

    widths = [16, 34, 16, 34, 10, 8, 11, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(HEADERS)), ws.max_row)
    # 已预填 HS 的行标淡绿，便于用户知道哪些不用动
    green = PatternFill('solid', fgColor='E2EFDA')
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=3).value:
            ws.cell(row=r, column=3).fill = green

    wb.save(OUT)
    filled = sum(1 for r in rows if r[2])
    print('saved %s | %d 行 | 已预填HS %d 行' % (OUT, len(rows), filled))


if __name__ == '__main__':
    main()
